import pytest
from xunit_rag.ingestion import DocumentIngestor, scan_docs
from xunit_rag.utils import file_sha


def test_scan_changes_and_original_preservation(config, tokenizer):
    sub = config.docs / "sub"
    sub.mkdir()
    text = sub / "Source.TXT"
    text.write_text("Original\n\nMore content", encoding="utf-8")
    unsupported = config.docs / "binary.xyz"
    unsupported.write_bytes(b"\x00\x01")
    before = file_sha(text)
    ingestor = DocumentIngestor(config, tokenizer)
    files = scan_docs(config)
    assert len(files) == 2 and next(f for f in files if f["path"] == "sub/Source.TXT")["supported"]
    assert ingestor.changes()["added"] == ["binary.xyz", "sub/Source.TXT"]
    leaves, manifest = ingestor.ingest()
    assert file_sha(text) == before and unsupported.read_bytes() == b"\x00\x01"
    assert manifest["files"]["binary.xyz"]["status"] == "unsupported"
    assert leaves and all(not leaf.pdf_pages and leaf.printed_page is None for leaf in leaves)
    assert leaves[0].location["text_lines"]
    again, _ = ingestor.ingest()
    assert [leaf.leaf_id for leaf in leaves] == [leaf.leaf_id for leaf in again]
    text.write_text("Changed source", encoding="utf-8")
    assert ingestor.changes()["changed"] == ["sub/Source.TXT"]
    text.unlink()
    assert ingestor.changes()["removed"] == ["sub/Source.TXT"]
    final, _ = ingestor.ingest()
    assert not final


def test_empty_docs_does_not_invent_corpus(config, tokenizer):
    leaves, manifest = DocumentIngestor(config, tokenizer).ingest()
    assert not leaves and manifest["leaf_count"] == 0
    assert list(config.docs.iterdir()) == []


def test_invalid_utf8_is_reported_not_silently_lost(config, tokenizer):
    (config.docs / "invalid.txt").write_bytes(b"\xff\xfe\x00")
    with pytest.raises(RuntimeError, match="Ingestão incompleta"):
        DocumentIngestor(config, tokenizer).ingest()


def test_ocr_skips_pages_with_native_text():
    from types import SimpleNamespace
    from xunit_rag.ingestion import SelectiveOcr

    calls = []

    def ocr(conversion, pages):
        calls.extend(page.page_no for page in pages)
        return iter(pages)

    pages = [SimpleNamespace(page_no=i) for i in range(1, 7)]
    assert list(SelectiveOcr(ocr, [1, 5])(None, pages)) == pages
    assert calls == [1, 5]


def test_prepare_empty_deactivates_old_index_without_inference(config, tokenizer, monkeypatch):
    from xunit_rag.cli import prepare
    from xunit_rag.storage import IndexStore
    from xunit_rag.utils import atomic_json, read_json

    monkeypatch.setattr("xunit_rag.cli.QwenTokenizer", lambda *args: tokenizer)
    pointer = config.data / "manifests" / "active.json"
    atomic_json(pointer, {"build_id": "historical"})
    result = prepare(config)
    assert result["status"] == "no_documents"
    assert read_json(pointer)["status"] == "unavailable"
    with pytest.raises(ValueError, match="Sem documentos"):
        IndexStore(config).active()


@pytest.mark.parametrize("format", ["md", "docx"])
def test_structured_text_formats(config, tokenizer, format):
    path = config.docs / f"source.{format}"
    if format == "md":
        path.write_text(
            "# State Verification\n\nInspect the resulting state after exercising the system.",
            encoding="utf-8",
        )
    else:
        from docx import Document

        document = Document()
        document.add_heading("State Verification", level=1)
        document.add_paragraph("Inspect the resulting state after exercising the system.")
        document.save(path)
    leaves, _ = DocumentIngestor(config, tokenizer).ingest()
    assert leaves and all(not leaf.pdf_pages for leaf in leaves)
    assert any("State Verification" in leaf.embedding_text for leaf in leaves)


def test_structured_conversion_checkpoint_survives_chunk_retry(config, tokenizer, monkeypatch):
    path = config.docs / "retry.md"
    path.write_text("# Test Stub\n\nProvide a controlled answer for a dependency.", encoding="utf-8")
    ingestor = DocumentIngestor(config, tokenizer)
    leaves, _ = ingestor.ingest()
    directory = config.data / "processed" / "documents" / leaves[0].doc_id
    (directory / "leaves.jsonl").unlink()
    from docling.document_converter import DocumentConverter

    monkeypatch.setattr(
        DocumentConverter,
        "convert",
        lambda *a, **k: (_ for _ in ()).throw(AssertionError("Must reuse conversion")),
    )
    again, _ = ingestor.ingest()
    assert [leaf.leaf_id for leaf in leaves] == [leaf.leaf_id for leaf in again]
